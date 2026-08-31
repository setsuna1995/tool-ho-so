# migrate_clean_and_tokenize_all_templates.py
import sys
from pathlib import Path
import docx
import word_writer

if sys.stdout:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

ROOT = Path(__file__).resolve().parent

DAO_DUC = ROOT / "01. Hồ sơ đạo đức đề cương - MẪU"
KHOA_HOC = ROOT / "02. Hồ sơ khoa học đề cương - MẪU"
MOI_CHUYEN_GIA = ROOT / "03. Công văn mời chuyên gia - MẪU"
NGHIEM_THU = ROOT / "04. Hồ sơ nghiệm thu - MẪU"


def migrate_00_qd_giao_de_tai(session: word_writer.Session) -> None:
    path = DAO_DUC / "00. QĐ Giao đề tài.docx"
    doc = session.open(path)
    # Bảng 3: ô Chủ nhiệm (hàng 2, cột 3) và ô Thành viên (hàng 3, cột 3)
    session.set_cell(doc, 3, 2, 3, "Chủ nhiệm đề tài:\r{{CHU_NHIEM_HO_TEN}} - {{CHU_NHIEM_DON_VI}}.")
    session.set_cell(doc, 3, 3, 3, "Thành viên thực hiện:\r{{DANH_SACH_NGHIEN_CUU_VIEN}}")
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_01_qdtlhd_dao_duc(session: word_writer.Session) -> None:
    path = DAO_DUC / "01. QĐTLHĐ đạo đức đề cương.docx"
    doc = session.open(path)
    # Table 2: 5 rows committee members
    placeholders = [
        ("[Họ và tên Chủ tịch HĐ]", "[Đơn vị công tác]"),
        ("[Họ và tên Phản biện 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Phản biện 2]", "[Đơn vị công tác]"),
        ("[Họ và tên Ủy viên 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Ủy viên 2]", "[Đơn vị công tác]"),
    ]
    for row_idx, (name_ph, org_ph) in enumerate(placeholders, start=1):
        session.set_cell(doc, 2, row_idx, 1, name_ph)
        session.set_cell(doc, 2, row_idx, 2, org_ph)
    # Table 3: 2 rows secretaries
    sec_placeholders = [
        ("[Họ và tên Thư ký 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Thư ký 2]", "[Đơn vị công tác]"),
    ]
    for row_idx, (name_ph, org_ph) in enumerate(sec_placeholders, start=1):
        session.set_cell(doc, 3, row_idx, 1, name_ph)
        session.set_cell(doc, 3, row_idx, 2, org_ph)
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_02_bb_hop_hd_dao_duc(session: word_writer.Session) -> None:
    path = DAO_DUC / "02. BB họp HĐ đạo đức.docx"
    doc = session.open(path)
    session.replace_text(
        doc,
        "Căn cứ Quyết định số: 04/QĐ-YHUD/2024 ngày 19 tháng 04 năm 2024 của Viện Y học ứng dụng Việt Nam",
        "Căn cứ Quyết định số: ……/QĐ-YHUD/{{NAM}} ngày …… tháng …… năm {{NAM}} của {{DON_VI_CHU_TRI}}",
    )
    session.replace_text(
        doc,
        "- Thời gian: ngày 25 tháng 04 năm 2024",
        "- Thời gian: ngày …… tháng …… năm {{NAM}}",
    )
    session.replace_text(
        doc,
        "- Chủ tịch Hội đồng: PGs. Ts. Hoàng Thị Thanh",
        "- Chủ tịch Hội đồng: {{CHU_TICH_HD_DAO_DUC}}",
    )
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_04_qd_chap_nhan_dao_duc(session: word_writer.Session) -> None:
    path = DAO_DUC / "04. QĐ chấp nhận đạo đức.docx"
    doc = session.open(path)
    session.replace_text(
        doc,
        "PGs. Ts. Hoàng Thị Thanh",
        "{{CHU_TICH_HD_DAO_DUC}}",
    )
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_05_qdtlhd_khoa_hoc(session: word_writer.Session) -> None:
    path = KHOA_HOC / "05. QĐ TLHĐ khoa học xét đề cương.docx"
    doc = session.open(path)
    # Table 2: 5 rows committee members
    placeholders = [
        ("[Họ và tên Chủ tịch HĐ]", "[Đơn vị công tác]"),
        ("[Họ và tên Phản biện 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Phản biện 2]", "[Đơn vị công tác]"),
        ("[Họ và tên Ủy viên 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Ủy viên 2]", "[Đơn vị công tác]"),
    ]
    for row_idx, (name_ph, org_ph) in enumerate(placeholders, start=1):
        session.set_cell(doc, 2, row_idx, 1, name_ph)
        session.set_cell(doc, 2, row_idx, 2, org_ph)
    # Table 3: 2 rows secretaries (cols 2 & 3)
    sec_placeholders = [
        ("[Họ và tên Thư ký 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Thư ký 2]", "[Đơn vị công tác]"),
    ]
    for row_idx, (name_ph, org_ph) in enumerate(sec_placeholders, start=1):
        session.set_cell(doc, 3, row_idx, 2, name_ph)
        session.set_cell(doc, 3, row_idx, 3, org_ph)
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_06_bb_hop_thong_qua_de_cuong(session: word_writer.Session) -> None:
    path = KHOA_HOC / "06. BB họp thông qua đề cương.docx"
    doc = session.open(path)
    session.replace_text(
        doc,
        "PGs. Ts. Hoàng Thị Thanh - Chủ tịch Hội đồng điều khiển phiên họp",
        "{{CHU_TICH_HD_KHOA_HOC}} - Chủ tịch Hội đồng điều khiển phiên họp",
    )
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_09_qd_thanh_lap_hd_nghiem_thu(session: word_writer.Session) -> None:
    path = NGHIEM_THU / "9. Quyết định thành lập HĐ nghiệm thu.docx"
    doc = session.open(path)
    # Table 2: 5 rows committee members (cols 2 & 3)
    placeholders = [
        ("[Họ và tên Chủ tịch HĐ]", "[Đơn vị công tác]"),
        ("[Họ và tên Phản biện 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Phản biện 2]", "[Đơn vị công tác]"),
        ("[Họ và tên Ủy viên 1]", "[Đơn vị công tác]"),
        ("[Họ và tên Ủy viên 2]", "[Đơn vị công tác]"),
    ]
    for row_idx, (name_ph, org_ph) in enumerate(placeholders, start=1):
        session.set_cell(doc, 2, row_idx, 2, name_ph)
        session.set_cell(doc, 2, row_idx, 3, org_ph)
    # Table 3: 2 rows secretaries (cols 1 & 2)
    sec_placeholders = [
        ("1. [Họ và tên Thư ký 1]", "[Đơn vị công tác]"),
        ("2. [Họ và tên Thư ký 2]", "[Đơn vị công tác]"),
    ]
    for row_idx, (name_ph, org_ph) in enumerate(sec_placeholders, start=1):
        session.set_cell(doc, 3, row_idx, 1, name_ph)
        session.set_cell(doc, 3, row_idx, 2, org_ph)
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_10_bb_hop_hd_nghiem_thu(session: word_writer.Session) -> None:
    path = NGHIEM_THU / "10. Biên bản họp HĐ nghiệm thu.docx"
    doc = session.open(path)
    session.replace_text(
        doc,
        "Tên 3 - Chủ tịch Hội đồng điều khiển phiên họp",
        "{{CHU_TICH_HD_NGHIEM_THU_TEN}} - Chủ tịch Hội đồng điều khiển phiên họp",
    )
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_phieu_cham_diem_tnls(session: word_writer.Session) -> None:
    path = NGHIEM_THU / "Phiếu chấm điểm nghiệm thu (TNLS).docx"
    doc = session.open(path)
    session.replace_text(
        doc,
        "Đồng chủ nhiệm: Tên 2",
        "Đồng chủ nhiệm: {{DONG_CHU_NHIEM_TEN}}",
    )
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_phieu_ky_nhan_tien(session: word_writer.Session) -> None:
    path = NGHIEM_THU / "Phiếu ký nhận tiền.docx"
    doc = session.open(path)
    # Table 2: rows 2 to 7, col 2
    placeholders = [
        "[Họ và tên Chủ tịch HĐ]",
        "[Họ và tên Phản biện 1]",
        "[Họ và tên Phản biện 2]",
        "[Họ và tên Ủy viên 1]",
        "[Họ và tên Ủy viên 2]",
        "[Họ và tên Thư ký HĐ]",
    ]
    for row_idx, name_ph in enumerate(placeholders, start=2):
        session.set_cell(doc, 2, row_idx, 2, name_ph)
    session.save_close(doc)
    print(f"Da migrate: {path.name}")


def migrate_all_templates() -> None:
    session = word_writer.Session()
    try:
        migrate_00_qd_giao_de_tai(session)
        migrate_01_qdtlhd_dao_duc(session)
        migrate_02_bb_hop_hd_dao_duc(session)
        migrate_04_qd_chap_nhan_dao_duc(session)
        migrate_05_qdtlhd_khoa_hoc(session)
        migrate_06_bb_hop_thong_qua_de_cuong(session)
        migrate_09_qd_thanh_lap_hd_nghiem_thu(session)
        migrate_10_bb_hop_hd_nghiem_thu(session)
        migrate_phieu_cham_diem_tnls(session)
        migrate_phieu_ky_nhan_tien(session)
    finally:
        session.quit()


if __name__ == "__main__":
    migrate_all_templates()
    print("XONG: Da token hoa va lam sach toan bo file mau.")
